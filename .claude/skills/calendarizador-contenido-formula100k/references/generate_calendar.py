#!/usr/bin/env python3
"""
Genera CSV y/o XLSX para un calendario de contenido FÓRMULA 100K.

Soporta DOS TIPOS DE PIEZA:
  - reels (default)      → columnas DÍA · IDEA · GUION · LLAMADO A LA ACCIÓN · FORMATO · REFERENCIA
  - carruseles            → columnas DÍA · FECHA · # · IDEA · GUION (SLIDE x SLIDE) · CTA · FORMATO
                            con altura de fila dinámica + coloreo por semana

Uso:
  python3 generate_calendar.py --tipo-pieza reels --modo crecimiento \
      --output-dir /path/ --data-json '{...}'

  python3 generate_calendar.py --tipo-pieza carruseles --modo crecimiento \
      --output-dir /path/ --output-name 00_CALENDARIO_CARRUSELES --data-json '{...}'

Input JSON (reels):
  {
    "modo": "crecimiento",
    "fecha_inicio": "2026-05-04",
    "filas": [
      {"dia":"Lunes","idea":"...","guion":"...","cta":"...","formato":"Pizarra (#19)","referencia":"Original"},
      ...
    ]
  }

Input JSON (carruseles):
  {
    "modo": "crecimiento",
    "fecha_inicio": "2026-05-18",
    "titulo": "CALENDARIO DE CARRUSELES — SEMANA 18 MAYO 2026",
    "subtitulo": "[TU NOMBRE] · [TU MARCA] · @tuhandle",
    "filas": [
      {"dia":"LUN","fecha":"18 MAY 2026","numero":"01","idea":"...","guion":"SLIDE 1 — ...\n\n...","cta":"MÁQUINA","formato":"Carrusel · 7 slides","semana":"S1"},
      ...
    ]
  }
"""

import argparse
import csv
import json
import os
import sys
from pathlib import Path

REELS_COLUMNS = ["DÍA", "IDEA", "GUION", "LLAMADO A LA ACCIÓN", "FORMATO", "REFERENCIA"]
CARRUSEL_COLUMNS = ["DÍA", "FECHA", "#", "IDEA", "GUION (SLIDE x SLIDE)", "CTA", "FORMATO"]


# ─────────────────────────────────────────────────────────────────────────────
# MODO REELS
# ─────────────────────────────────────────────────────────────────────────────

def write_csv_reels(path: Path, filas: list[dict]) -> None:
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f, quoting=csv.QUOTE_ALL)
        writer.writerow(REELS_COLUMNS)
        for row in filas:
            writer.writerow([
                row.get("dia", ""),
                row.get("idea", ""),
                row.get("guion", ""),
                row.get("cta", ""),
                row.get("formato", ""),
                row.get("referencia", "Original"),
            ])


def write_xlsx_reels(path: Path, filas: list[dict], modo: str) -> bool:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        return False

    wb = Workbook()
    ws = wb.active
    ws.title = f"Calendario {modo[:25]}"

    header_fill = PatternFill("solid", fgColor="44403C")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    cell_align = Alignment(wrap_text=True, vertical="top")
    thin = Side(border_style="thin", color="D6D3D1")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col_idx, col_name in enumerate(REELS_COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    type_colors = {
        "Viral": "FEE2E2",
        "Valor": "FEF3C7",
        "Venta": "D1FAE5",
    }

    for row_idx, row in enumerate(filas, start=2):
        values = [
            row.get("dia", ""),
            row.get("idea", ""),
            row.get("guion", ""),
            row.get("cta", ""),
            row.get("formato", ""),
            row.get("referencia", "Original"),
        ]
        tipo = row.get("tipo", "")
        fill = PatternFill("solid", fgColor=type_colors.get(tipo, "FFFFFF")) if tipo in type_colors else None
        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = cell_align
            cell.border = border
            if fill and col_idx == 1:
                cell.fill = fill
                cell.font = Font(bold=True)

    widths = [12, 28, 70, 32, 24, 38]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    for row_idx, row in enumerate(filas, start=2):
        guion = row.get("guion", "")
        line_count = guion.count("\n") + 1
        ws.row_dimensions[row_idx].height = max(min(line_count * 15, 460), 90)

    ws.freeze_panes = "A2"
    wb.save(path)
    return True


# ─────────────────────────────────────────────────────────────────────────────
# MODO CARRUSELES
# ─────────────────────────────────────────────────────────────────────────────

def write_xlsx_carruseles(
    path: Path,
    filas: list[dict],
    titulo: str = "",
    subtitulo: str = "",
) -> bool:
    """Genera XLSX para tipo carruseles.

    Layout:
    - Fila 1: título mergeado A1:G1
    - Fila 2: subtítulo mergeado A2:G2
    - Fila 4: headers
    - Fila 5+: datos (una fila por carrusel)
    - Coloreo DÍA por semana (S1 amarillo, S2 lila)
    - Altura dinámica por celda de guion
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    except ImportError:
        return False

    wb = Workbook()
    ws = wb.active
    ws.title = "Calendario Carruseles"

    # Estilos
    header_fill = PatternFill("solid", fgColor="1A1A1A")
    header_font = Font(name="Calibri", size=12, bold=True, color="FCE96B")
    title_font = Font(name="Calibri", size=16, bold=True, color="1A1A1A")
    subtitle_font = Font(name="Calibri", size=11, italic=True, color="3A3A3A")
    cell_font = Font(name="Calibri", size=10, color="1A1A1A")
    guion_font = Font(name="Calibri", size=10, color="1A1A1A")
    cta_font = Font(name="Calibri", size=11, bold=True, color="1A1A1A")
    week1_fill = PatternFill("solid", fgColor="FCE96B")
    week2_fill = PatternFill("solid", fgColor="D9CCEC")
    even_fill = PatternFill("solid", fgColor="F5EFE3")
    odd_fill = PatternFill("solid", fgColor="FFFFFF")
    thin = Side(border_style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # Título + subtítulo
    if titulo:
        ws.merge_cells("A1:G1")
        ws["A1"] = titulo
        ws["A1"].font = title_font
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 28
    if subtitulo:
        ws.merge_cells("A2:G2")
        ws["A2"] = subtitulo
        ws["A2"].font = subtitle_font
        ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[2].height = 22

    # Headers (fila 4)
    for col_idx, h in enumerate(CARRUSEL_COLUMNS, start=1):
        cell = ws.cell(row=4, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    ws.row_dimensions[4].height = 30

    # Filas de datos
    current_row = 5
    for row in filas:
        dia = row.get("dia", "")
        fecha = row.get("fecha", "")
        numero = row.get("numero", "")
        idea = row.get("idea", "")
        guion = row.get("guion", "")
        cta = row.get("cta", "")
        formato = row.get("formato", "")
        semana = row.get("semana", "S1")

        fill = week1_fill if semana == "S1" else week2_fill
        row_fill = even_fill if current_row % 2 == 0 else odd_fill

        values = [dia, fecha, numero, idea, guion, cta, formato]
        for col_idx, val in enumerate(values, start=1):
            cell = ws.cell(row=current_row, column=col_idx, value=val)

            if col_idx == 5:  # GUION
                cell.font = guion_font
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            elif col_idx == 6:  # CTA
                cell.font = cta_font
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            elif col_idx == 4:  # IDEA
                cell.font = cell_font
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            else:
                cell.font = cell_font
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

            cell.border = border

            if col_idx == 1:
                cell.fill = fill
                cell.font = Font(name="Calibri", size=11, bold=True, color="1A1A1A")
            else:
                cell.fill = row_fill

        # Altura dinámica según líneas en guion
        line_count = guion.count("\n") + 1
        ws.row_dimensions[current_row].height = max(line_count * 14, 200)
        current_row += 1

    # Anchos
    widths = {"A": 8, "B": 14, "C": 5, "D": 32, "E": 90, "F": 14, "G": 22}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    ws.freeze_panes = "A5"
    wb.save(path)
    return True


# ─────────────────────────────────────────────────────────────────────────────
# MAIN
# ─────────────────────────────────────────────────────────────────────────────

def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--tipo-pieza", choices=["reels", "carruseles"], default="reels")
    parser.add_argument("--modo", required=True)
    parser.add_argument("--output-dir", required=True)
    parser.add_argument("--output-name", required=False, help="Nombre base del archivo sin extensión. Si no se pasa, se autogenera.")
    parser.add_argument("--data-json", required=False)
    parser.add_argument("--data-file", required=False)
    args = parser.parse_args()

    if args.data_file:
        with open(args.data_file, "r", encoding="utf-8") as f:
            data = json.load(f)
    elif args.data_json:
        data = json.loads(args.data_json)
    else:
        data = json.load(sys.stdin)

    modo = data.get("modo", args.modo)
    fecha = data.get("fecha_inicio", "")
    filas = data.get("filas", [])

    if not filas:
        print("ERROR: No filas to write.", file=sys.stderr)
        return 1

    out_dir = Path(args.output_dir)
    out_dir.mkdir(parents=True, exist_ok=True)

    if args.output_name:
        base = args.output_name
    else:
        base = f"calendario_{modo}_{fecha}" if fecha else f"calendario_{modo}"

    if args.tipo_pieza == "reels":
        csv_path = out_dir / f"{base}.csv"
        xlsx_path = out_dir / f"{base}.xlsx"
        write_csv_reels(csv_path, filas)
        xlsx_ok = write_xlsx_reels(xlsx_path, filas, modo)
        result = {
            "tipo_pieza": "reels",
            "csv": str(csv_path),
            "xlsx": str(xlsx_path) if xlsx_ok else None,
            "xlsx_error": None if xlsx_ok else "openpyxl no instalado. Instala con: pip3 install openpyxl",
            "rows": len(filas),
        }
    else:  # carruseles
        xlsx_path = out_dir / f"{base}.xlsx"
        titulo = data.get("titulo", "")
        subtitulo = data.get("subtitulo", "")
        xlsx_ok = write_xlsx_carruseles(xlsx_path, filas, titulo=titulo, subtitulo=subtitulo)
        result = {
            "tipo_pieza": "carruseles",
            "csv": None,
            "xlsx": str(xlsx_path) if xlsx_ok else None,
            "xlsx_error": None if xlsx_ok else "openpyxl no instalado. Instala con: pip3 install openpyxl",
            "rows": len(filas),
        }

    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    sys.exit(main())
